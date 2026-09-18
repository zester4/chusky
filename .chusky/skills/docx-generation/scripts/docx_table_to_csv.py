#!/usr/bin/env python3
import argparse, csv
from docx import Document

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--table", type=int, default=0)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    doc = Document(args.docx)
    if args.table >= len(doc.tables):
        raise SystemExit("Table index out of range")
    with open(args.out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        for row in doc.tables[args.table].rows:
            w.writerow([cell.text for cell in row.cells])
    print(args.out)

if __name__ == "__main__":
    main()
