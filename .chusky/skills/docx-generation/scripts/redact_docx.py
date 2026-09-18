#!/usr/bin/env python3
import argparse, re
from pathlib import Path

EMAIL = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
PHONE = re.compile(r"(?<!\w)(?:\+?\d[\d .()\-]{7,}\d)(?!\w)")

def replace_runs(paragraph, patterns, replacement):
    for run in paragraph.runs:
        txt = run.text
        for pat in patterns:
            txt = pat.sub(replacement, txt)
        run.text = txt

def walk_table(table, patterns, replacement):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                replace_runs(p, patterns, replacement)
            for t in cell.tables:
                walk_table(t, patterns, replacement)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--emails", action="store_true")
    ap.add_argument("--phones", action="store_true")
    ap.add_argument("--literal", action="append", default=[])
    ap.add_argument("--replacement", default="[REDACTED]")
    args = ap.parse_args()
    try:
        from docx import Document
    except ImportError:
        raise SystemExit("python-docx is required: pip install python-docx")
    pats = []
    if args.emails: pats.append(EMAIL)
    if args.phones: pats.append(PHONE)
    pats.extend(re.compile(re.escape(x), re.I) for x in args.literal)
    doc = Document(args.input)
    for p in doc.paragraphs:
        replace_runs(p, pats, args.replacement)
    for t in doc.tables:
        walk_table(t, pats, args.replacement)
    for sec in doc.sections:
        for container in (sec.header, sec.footer):
            for p in container.paragraphs:
                replace_runs(p, pats, args.replacement)
            for t in container.tables:
                walk_table(t, pats, args.replacement)
    doc.save(args.output)
    print(args.output)

if __name__ == "__main__":
    main()
