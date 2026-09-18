#!/usr/bin/env python3
import argparse, collections, json

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    args = ap.parse_args()
    try:
        from docx import Document
    except ImportError:
        raise SystemExit("python-docx is required: pip install python-docx")
    doc = Document(args.docx)
    para_styles = collections.Counter(p.style.name if p.style else "(none)" for p in doc.paragraphs)
    direct = 0
    fonts = collections.Counter()
    sizes = collections.Counter()
    for p in doc.paragraphs:
        for r in p.runs:
            if r.bold is not None or r.italic is not None or r.underline is not None or r.font.name or r.font.size:
                direct += 1
            if r.font.name:
                fonts[r.font.name] += 1
            if r.font.size:
                sizes[str(round(r.font.size.pt, 2))] += 1
    out = {
        "paragraph_styles": para_styles,
        "directly_formatted_runs": direct,
        "explicit_fonts": fonts,
        "explicit_sizes_pt": sizes
    }
    print(json.dumps(out, indent=2, default=dict))

if __name__ == "__main__":
    main()
