#!/usr/bin/env python3
import argparse
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    wb = load_workbook(args.xlsx, data_only=True, read_only=True)
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]
    rows = [[("" if v is None else str(v)) for v in row] for row in ws.iter_rows(values_only=True)]
    if not rows:
        raise SystemExit("Worksheet is empty")
    doc = Document()
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for j, value in enumerate(rows[0]):
        table.rows[0].cells[j].text = value
    repeat_header(table.rows[0])
    for row in rows[1:]:
        cells = table.add_row().cells
        for j in range(len(cells)):
            cells[j].text = row[j] if j < len(row) else ""
    doc.save(args.out)
    print(args.out)

if __name__ == "__main__":
    main()
